import json, re, urllib.request, io
def raw(u):
    req = urllib.request.Request(u, headers={'User-Agent': 'Mozilla/5.0 MonitoraBrasil', 'Accept': '*/*'})
    with urllib.request.urlopen(req, timeout=90) as r: return r.read()
def get(u, n=2500):
    try:
        b = raw(u); print(f'\n=== {u}\nbytes {len(b)}'); print(b[:n].decode('utf-8','replace')); return b
    except Exception as e: print(f'\n=== {u}\nERRO {e}')
# 1) Obrasgov: rotas citadas na página de documentação
h = raw('https://api-publica.obrasgov.gestao.gov.br/').decode('utf-8','replace')
rotas = sorted(set(re.findall(r'["\'(](/[a-zA-Z0-9_\-/{}]*(?:projeto|execucao|historico|situacao|api)[a-zA-Z0-9_\-/{}]*)', h)))
print('ROTAS', rotas[:80])
urls = sorted(set(re.findall(r'https?://[^"\'\s<>]+', h)))
print('URLS', [u for u in urls if 'obrasgov' in u or 'api' in u][:60])
i = h.find('projeto-investimento'); print('TRECHO', h[max(0,i-1500):i+2500] if i>=0 else 'sem projeto-investimento')
for p in ['/obrasgov/api/projeto-investimento?pagina=0&tamanhoDaPagina=2', '/projeto-investimento?pagina=0&tamanhoDaPagina=2', '/api/projeto-investimento?pagina=0&tamanhoDaPagina=2', '/obrasgov/projeto-investimento?pagina=0&tamanhoDaPagina=2']:
    get('https://api-publica.obrasgov.gestao.gov.br'+p, 1500)
# 2) Tesouro RTN
import subprocess; subprocess.run(['pip','install','-q','openpyxl'])
import openpyxl
x = raw('https://www.tesourotransparente.gov.br/ckan/dataset/ab56485b-9c40-4efb-8563-9ce3e1973c4b/resource/527ccdb1-3059-42f3-bf23-b5e3ab4c6dc6/download/seriehistoricajul26.xlsx')
wb = openpyxl.load_workbook(io.BytesIO(x), read_only=True, data_only=True)
print('ABAS', wb.sheetnames)
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    print(f'\n--- ABA {ws.title}: {len(rows)} linhas x {max((len(r) for r in rows), default=0)} col')
    for k, r in enumerate(rows[:8]): print(k, [c for c in r[:8]])
    for k, r in enumerate(rows):
        t = ' '.join(str(c) for c in r[:3] if c is not None)
        if re.search(r'nvestimento|PIB|Produto Interno', t): print('  ACHOU', k, [c for c in r[:6]], '...', [c for c in r[-4:]])
get('https://apiapex.tesouro.gov.br/aria/v1/series-temporais/docs', 3000)
